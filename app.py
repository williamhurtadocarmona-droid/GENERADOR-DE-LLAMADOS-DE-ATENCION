import streamlit as st
import pandas as pd
from docx import Document
import io

# Configuración inicial de la página
st.set_page_config(
    page_title="Generador de Documentos SENA",
    page_icon="📄",
    layout="centered"
)

st.title("📄 Generador de Documentos Automatizado")
st.write("Sube tu archivo de Excel, selecciona el aprendiz y genera el archivo de Word personalizado.")

# 1. Cargar el archivo de Excel desde la interfaz web
excel_file = st.file_uploader("1. Sube tu archivo de Excel (.xlsx)", type=["xlsx", "xls"])

if excel_file is not None:
    # Leer todas las hojas o la primera por defecto
    df = pd.read_excel(excel_file)
    st.success("¡Archivo de Excel cargado correctamente!")
    
    # 2. Selección de la columna que identifica al Aprendiz
    columnas = list(df.columns)
    col_aprendiz = st.selectbox("2. Selecciona la columna que tiene los nombres de los Aprendices:", columnas)
    
    if col_aprendiz:
        # Filtrar lista única de aprendices eliminando valores vacíos
        lista_aprendices = df[col_aprendiz].dropna().unique()
        
        # 3. Menú desplegable para seleccionar un Aprendiz específico
        aprendiz_seleccionado = st.selectbox("3. Selecciona el Aprendiz a procesar:", lista_aprendices)
        
        if aprendiz_seleccionado:
            # Obtener la fila de datos correspondiente al aprendiz seleccionado
            fila_datos = df[df[col_aprendiz] == aprendiz_seleccionado].iloc[0]
            
            # Mostrar los datos detectados en una tarjeta desplegable
            with st.expander("Ver datos detectados de este registro"):
                for col in columnas:
                    st.write(f"**{col}:** {fila_datos[col]}")
            
            # 4. Botón para procesar y generar el documento
            if st.button("🚀 Generar Documento de Word"):
                try:
                    # Abrir la plantilla de Word cargada en el repositorio
                    doc = Document("plantilla.docx")
                    
                    # Crear diccionario dinámico con todas las columnas del Excel
                    # Reemplazará {{NOMBRE_COLUMNA}} por el valor de la fila
                    reemplazos = {f"{{{{{col}}}}}" : str(fila_datos[col]) for col in columnas}
                    
                    # Reemplazar texto en los párrafos normales
                    for p in doc.paragraphs:
                        for tag, valor in reemplazos.items():
                            if tag in p.text:
                                p.text = p.text.replace(tag, valor)
                    
                    # Reemplazar texto dentro de las tablas (si la plantilla contiene tablas)
                    for tabla in doc.tables:
                        for fila_tabla in tabla.rows:
                            for celda in fila_tabla.cells:
                                for p in celda.paragraphs:
                                    for tag, valor in reemplazos.items():
                                        if tag in p.text:
                                            p.text = p.text.replace(tag, valor)
                    
                    # Guardar el documento en la memoria temporal (Buffer)
                    output = io.BytesIO()
                    doc.save(output)
                    output.seek(0)
                    
                    # Botón de descarga para el usuario
                    st.download_button(
                        label="📥 Descargar Documento Generado",
                        data=output,
                        file_name=f"Documento_{aprendiz_seleccionado}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    )
                    st.success("¡Documento generado con éxito!")
                    
                except FileNotFoundError:
                    st.error("Error: No se encontró el archivo 'plantilla.docx' en el repositorio de GitHub. Asegúrate de haberlo subido con ese nombre exacto.")